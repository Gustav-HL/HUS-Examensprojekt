from pathlib import Path
from pypdf import PdfReader
from docx import Document
import json

class FileHandler:
    def __init__(self):
        pass

    def extract_text(self, path) -> str:
        path = Path(path)

        suffix = path.suffix.lower()

        if suffix == ".pdf":
            reader = PdfReader(path)
            return "\n".join(page.extract_text() or " " for page in reader.pages)

        if suffix == ".docx":
            document = Document(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)

        if suffix == ".txt":
            with open(path, "r", encoding="utf-8") as file:
                return file.read()

        raise ValueError(f"'{suffix}' is not supported :(")

    def extract_metadata(self, path) -> dict: 
        path = Path(path)
        suffix = path.suffix.lower()
        data = {}

        if suffix == ".pdf":
            reader = PdfReader(path)
            metadata = reader.metadata  #fetches metadata

            if metadata:
                data["title"] = metadata.title
                data["author"] = metadata.author
                data["creator"] = metadata.creator
                data["creation_date"] = str(metadata.creation_date) if metadata.creation_date else None
                data["modification_date"] = str(metadata.modification_date) if metadata.modification_date else None

        elif suffix == ".docx":
            document = Document(str(path))
            metadata = document.core_properties  #fetches metadata

            data["title"] = metadata.title
            data["author"] = metadata.author
            data["creator"] = metadata.author
            data["creation_date"] = str(metadata.created) if metadata.created else None
            data["modification_date"] = str(metadata.modified) if metadata.modified else None

        elif suffix == ".txt":
            return {}

        else:
            raise ValueError(f"'{suffix}' is not supported :(")

        return data