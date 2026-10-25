from pathlib import Path
from pypdf import PdfReader
from docx import Document
import json

def extract_text(path) -> str:
    path = Path(path)

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or " " for page in reader.pages)

    if suffix == ".docx":
        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    if suffix == ".txt":
        with open(path, "r", encoding="utf-8") as file:
            return file.read()

    raise ValueError(f"'{suffix}' is not supported :(")

def extract_metadata(path) -> dict:
    path = Path(path)

    suffix = path.suffix.lower()

    data = {}

    if suffix == ".pdf":
        reader = PdfReader(path)
        metadata = reader.metadata #fetches metadata

        if metadata:
            data["title"] = metadata.title
            data["author"] = metadata.author
            data["creator"] = metadata.creator
            data["creation_date"] = str(metadata.creation_date)
            data["modification_date"] = str(metadata.modification_date)

    elif suffix == ".docx":
        document = Document(str(path))
        metadata = document.core_properties #fetches metadata
        
        data["title"] = metadata.title
        data["author"] = metadata.author
        data["creator"] = metadata.author
        data["creation_date"] = metadata.created
        data["modification_date"] = metadata.modified

    else:
        raise ValueError(f"'{suffix}' is not supported :(")
    
    return data

if __name__ == "__main__":
    folder = Path(__file__).parent
    path = folder / "riktlinjer-for-region-skanes-miljovardsfond-run-2024-03-08.pdf"

    print("TEXT:")
    print(extract_text(path))
    print("METADATA:")
    print(extract_metadata(path))